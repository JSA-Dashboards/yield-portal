"""
Checks for the Report text page's layout helpers and its Word / PDF downloads.

    python tests/test_report_text.py

The reports below are MADE UP (the repo is public).
"""
import datetime as dt
import io
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import pandas as pd  # noqa: E402

import report_text as RT  # noqa: E402

failures = []


def check(label, got, want):
    if got != want:
        failures.append(f"{label}: got {got!r}, want {want!r}")


# --- the place label ---------------------------------------------------------------
check("label up to the colon",
      RT.label_and_body("Adams Co, IN: 248 bpa, 18% moisture.", "Adams Co", "IN"),
      ("Adams Co, IN:", "248 bpa, 18% moisture."))
check("no colon nearby: just the place",
      RT.label_and_body("Story Co IA 3 fields went 61-64 bpa", "Story Co", "IA"),
      ("Story Co", "IA 3 fields went 61-64 bpa"))
check("place only in the email subject: put in front",
      RT.label_and_body("80 acres went 96 bu/acre, a record.", "Clark Co", "IL"),
      ("Clark Co, IL:", "80 acres went 96 bu/acre, a record."))
check("place named later: left as written",
      RT.label_and_body("Beans in western Clark Co ran 70.", "Clark Co", "IL"),
      ("", "Beans in western Clark Co ran 70."))
check("no place", RT.label_and_body("Delta: beans averaged 70.", None, "MS"),
      ("", "Delta: beans averaged 70."))
check("NaN place", RT.label_and_body("Delta: beans averaged 70.", float("nan"), "MS"),
      ("", "Delta: beans averaged 70."))

# --- Markdown: report text can't turn into formatting ------------------------------
md = RT.md_report("Polk Co, IA: ~10% under LY, $4 corn, field #2 [dry] *ugh*", "Polk Co", "IA")
check("markdown escaped", md,
      r"**Polk Co, IA:** \~10% under LY, \$4 corn, field \#2 \[dry\] \*ugh\*")
check("search term highlighted in label and text",
      RT.md_report("Polk Co, IA: rust in Polk too", "Polk Co", "IA", "polk"),
      "**:yellow-background[Polk] Co, IA:** rust in :yellow-background[Polk] too")
check("regex characters in a search are literal",
      RT.md_report("Polk Co, IA: 200+ bpa", "Polk Co", "IA", "200+"),
      "**Polk Co, IA:** :yellow-background[200+] bpa")

# --- dates and order -----------------------------------------------------------------
check("day label", RT.day_label(dt.date(2026, 9, 4)), "Sep 4")
check("no date", [RT.day_label(v) for v in (None, float("nan"), pd.NaT)], ["", "", ""])
check("state names", [RT.state_name(s) for s in ("IA", "ND", None)],
      ["Iowa", "North Dakota", "State not given"])

df = pd.DataFrame([
    dict(crop="Soybeans", state="IA", location="Polk Co", date_reported=dt.date(2026, 9, 20),
         raw_text="Polk Co, IA: 62 bpa."),
    dict(crop="Corn", state="IL", location="Clark Co", date_reported=None,
         raw_text="Clark Co, IL: 230 bpa."),
    dict(crop="Corn", state="IL", location="Adams Co", date_reported=dt.date(2026, 9, 25),
         raw_text="Adams Co, IL: 241 bpa."),
    dict(crop="Corn", state="IA", location="Story Co", date_reported=dt.date(2026, 9, 28),
         raw_text="Story Co, IA: 255 bpa."),
    dict(crop="Corn", state="IL", location="Brown Co", date_reported=dt.date(2026, 9, 21),
         raw_text="Brown Co, IL: 222 bpa — “so far”…"),
])
check("PDF order: corn first, states by name, dated by date then undated",
      list(RT.ordered(df)["location"]), ["Brown Co", "Adams Co", "Clark Co", "Story Co", "Polk Co"])
check("place order", list(RT.ordered(df, by="place")["location"]),
      ["Adams Co", "Brown Co", "Clark Co", "Story Co", "Polk Co"])
check("sections", [(c, [n for n, _ in s]) for c, s in RT.sections(RT.ordered(df))],
      [("Corn", ["Illinois", "Iowa"]), ("Soybeans", ["Iowa"])])
halves = pd.DataFrame([
    dict(crop="Corn", state="NE", location="Saline Co", date_reported=dt.date(2026, 10, 5),
         raw_text="Dryland 150 bpa while irrigated 215 bpa.", irrigation=p) for p in
    ("Non-irrigated", "Irrigated")])
check("a report split by practice shows its words once", len(RT.ordered(halves)), 1)
check("...while different reports at one place both show",
      len(RT.ordered(pd.concat([halves, df.assign(location="Saline Co", state="NE")]))), 6)

table = RT.md_table(RT.ordered(df)[lambda d: d["state"] == "IL"], query="bpa")
check("table: header + one line per report", table.count("\n") + 1, 2 + 3)
check("table: dated row", table.splitlines()[2],
      "| :gray[Sep 21] | **Brown Co, IL:** 222 :yellow-background[bpa] — “so far”… |")
check("table: undated row has an empty date cell", table.splitlines()[4].startswith("|  | **Clark"),
      True)
check("table without dates", RT.md_table(df.iloc[[1]], dates=False).splitlines(),
      ["| Report |", "| --- |", "| **Clark Co, IL:** 230 bpa. |"])
check("a pipe in a report can't split the cell",
      RT.md_report("Polk Co, IA: 60|62 bpa", "Polk Co", "IA"), r"**Polk Co, IA:** 60\|62 bpa")
check("JSA's own reports are tagged; Ag Trader Talk's aren't",
      (RT.md_report("Polk Co, IA: 210 bpa", "Polk Co", "IA", source="JSA"),
       RT.md_report("Polk Co, IA: 210 bpa", "Polk Co", "IA", source="Ag Trader Talk")),
      ("**Polk Co, IA:** 210 bpa :blue-badge[JSA]", "**Polk Co, IA:** 210 bpa"))
check("a seed plot is tagged with its company",
      RT.md_report("Polk Co, IA: plot average 232 bpa", "Polk Co", "IA", source="Seed plot",
                   company="Beck's"),
      "**Polk Co, IA:** plot average 232 bpa :orange-badge[Beck's plot]")
check("tags", [RT.source_tag("Seed plot"), RT.source_tag("Ag Trader Talk", "x.pdf"),
               RT.source_tag("Seed plot", "Pioneer")], ["Seed plot", "", "Pioneer plot"])

# --- downloads -----------------------------------------------------------------------
pdf = RT.to_pdf(RT.ordered(df), "Yield reports · 2026", "made-up rows")
import fitz  # noqa: E402

doc = fitz.open("pdf", pdf)
text = "".join(p.get_text() for p in doc)
check("pdf: one page", doc.page_count, 1)
for want in ("Yield reports · 2026", "Corn", "Illinois", "Sep 21", "Brown Co, IL:",
             "222 bpa — “so far”…", "Page 1 of 1"):
    check(f"pdf has {want!r}", want in text, True)

from docx import Document  # noqa: E402

word = Document(io.BytesIO(RT.to_docx(RT.ordered(df), "Yield reports · 2026", "made-up rows")))
paras = [p.text for p in word.paragraphs]
check("word: headings", [p.text for p in word.paragraphs if p.style.name.startswith("Heading")],
      ["Corn", "Illinois", "Iowa", "Soybeans", "Iowa"])
check("word: a dated report", "Sep 21   Brown Co, IL: 222 bpa — “so far”…" in paras, True)
check("word: an undated report", "Clark Co, IL: 230 bpa." in paras, True)
bold = [r.text for p in word.paragraphs for r in p.runs if r.bold]
check("word: labels bold", "Adams Co, IL: " in bold, True)

# --- yellow: reported since the week the last Tuesday email covered ------------------
check("last report week, on a Thursday", RT.last_report_week(dt.date(2026, 10, 8)),
      (dt.date(2026, 9, 29), dt.date(2026, 10, 5)))
check("...on the Tuesday itself", RT.last_report_week(dt.date(2026, 10, 6)),
      (dt.date(2026, 9, 29), dt.date(2026, 10, 5)))
check("...on the Monday before the next", RT.last_report_week(dt.date(2026, 10, 12)),
      (dt.date(2026, 9, 29), dt.date(2026, 10, 5)))
since = dt.date(2026, 9, 24)
check("new: on or after the day; never undated",
      [RT.is_new(d, since) for d in (dt.date(2026, 9, 24), dt.date(2026, 9, 23), None,
                                     pd.Timestamp("2026-09-30"))], [True, False, False, True])
check("no day, nothing new", RT.is_new(dt.date(2026, 9, 30), None), False)
il = RT.ordered(df)[lambda d: d["state"] == "IL"]
check("page: a new report's date in yellow",
      [":yellow-background[Sep 25]" in RT.md_table(il, new_since=since),
       ":gray[Sep 21]" in RT.md_table(il, new_since=since)], [True, True])
check("pdf html: one yellow paragraph per new report (Adams 9/25, Story 9/28)",
      RT.report_html(RT.ordered(df), "t", "s", since).count(f"background-color:{RT.NEW_BG}"), 2)


def yellow_fills(pdf_bytes):
    d = fitz.open("pdf", pdf_bytes)
    return sum(1 for pg in d for g in pg.get_drawings()
               if g.get("fill") and abs(g["fill"][2] - 0.627) < 0.01 and g["fill"][0] > 0.99)


check("pdf: yellow behind the new reports, none without a day",
      (yellow_fills(RT.to_pdf(RT.ordered(df), "t", "s", since)) >= 2,
       yellow_fills(RT.to_pdf(RT.ordered(df), "t", "s"))), (True, 0))
from docx.enum.text import WD_COLOR_INDEX  # noqa: E402

hl = Document(io.BytesIO(RT.to_docx(RT.ordered(df), "t", "s", since)))
lit = sorted({p.text.split(":")[0].split("   ")[-1] for p in hl.paragraphs
              if p.runs and all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)})
check("word: the new reports highlighted", lit, ["Adams Co, IL", "Story Co, IA"])

if failures:
    print(f"{len(failures)} FAILED:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("All report-text checks pass.")
