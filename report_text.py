"""
The reports in their own words, laid out like the annual yield PDF: crop, then
state, then one paragraph per report with the date it was reported. Shared by
the Report text page and its Word and PDF downloads.

Report text is free text, so it is always escaped for the medium it goes into:
Markdown on the page (a "~10%" or "$395" must not turn into strikethrough or
math), HTML for the PDF.
"""
import datetime as _dt
import html
import io
import re

import pandas as pd

from parse_pdfs import STATE_ABBR

STATE_NAME = {ab: name.title() for name, ab in STATE_ABBR.items()}
CROP_ORDER = {"Corn": 0, "Soybeans": 1}
_MD_SPECIAL = re.compile(r"([\\`*_~$\[\]<>#|])")
NEW_BG = "#fff3a0"         # new reports, in yellow, as the weekly email marks them


def last_report_week(today):
    """(first, last) report dates the latest Tuesday email called new: the seven days
    before the most recent Tuesday on or before `today` (weekly_email.send_day/week).
    Report text highlights from the first of them, so a download shows what that
    email called new and everything reported since."""
    day = today - _dt.timedelta(days=(today.weekday() - 1) % 7)
    return day - _dt.timedelta(days=7), day - _dt.timedelta(days=1)


def is_new(d, since) -> bool:
    """Reported on or after `since` (a report with no date never is)."""
    if since is None or d is None or not isinstance(d, _dt.date) or pd.isna(d):
        return False
    return (d.date() if isinstance(d, _dt.datetime) else d) >= since


def state_name(ab) -> str:
    if not isinstance(ab, str) or not ab:
        return "State not given"
    return STATE_NAME.get(ab, ab)


def day_label(d) -> str:
    """'Sep 24' for a date; '' when the report has none."""
    if not isinstance(d, _dt.date) or pd.isna(d):
        return ""
    return f"{d:%b} {d.day}"


def label_and_body(raw, location, state):
    """Split a report into its place label and the rest.

    Most reports open with their place ("Adams Co, IN: 248 bpa ..."), which
    becomes the label up to its colon. A report that leaves its place to the
    email subject gets "Place, ST:" in front; one that names it later in the
    text, or has no place, is left as written with no label."""
    raw = (raw or "").strip()
    loc = location.strip() if isinstance(location, str) else ""
    if not loc:
        return "", raw
    if raw.lower().startswith(loc.lower()):
        colon = raw.find(":", len(loc), len(loc) + 20)
        cut = colon + 1 if colon >= 0 else len(loc)
        return raw[:cut], raw[cut:].strip()
    if loc.lower() in raw.lower():
        return "", raw
    return (f"{loc}, {state}:" if isinstance(state, str) and state else f"{loc}:"), raw


def ordered(df: pd.DataFrame, by: str = "date") -> pd.DataFrame:
    """Corn before soybeans, states by name, then by report date (undated after,
    by place) — the PDF's order — or by place. A report split by practice is two
    rows with one text (dryland and irrigated): its words show once."""
    keys = pd.DataFrame({
        "_crop": df["crop"].map(CROP_ORDER).fillna(len(CROP_ORDER)),
        "_state": df["state"].map(state_name),
        "_date": pd.to_datetime(df["date_reported"], errors="coerce"),
        "_place": df["location"].fillna("").str.lower(),
    }, index=df.index)
    cols = ["_crop", "_state"] + (["_place", "_date"] if by == "place" else ["_date", "_place"])
    out = df.loc[keys.sort_values(cols, na_position="last", kind="stable").index]
    return out[~out[["crop", "state", "location", "raw_text"]].astype(str).duplicated().values]


def sections(df: pd.DataFrame):
    """(crop, [(state name, rows), ...]) in the frame's order."""
    for crop, g in df.groupby("crop", sort=False):
        yield crop, list(g.groupby(g["state"].map(state_name), sort=False))


# --- Markdown (the page) ------------------------------------------------------
def md_escape(s: str) -> str:
    return _MD_SPECIAL.sub(r"\\\1", s)


def _highlight(text: str, query: str) -> str:
    if not query:
        return md_escape(text)
    parts = re.split(f"({re.escape(query)})", text, flags=re.I)
    return "".join(f":yellow-background[{md_escape(p)}]" if i % 2 else md_escape(p)
                   for i, p in enumerate(parts) if p)


def source_tag(source, company=None):
    """The tag a report carries beside its text: 'JSA', '<company> plot', or ''
    for Ag Trader Talk's (most of them)."""
    if source == "JSA":
        return "JSA"
    if source == "Seed plot":
        return f"{company} plot" if isinstance(company, str) and company else "Seed plot"
    return ""


def md_report(raw, location, state, query: str = "", source=None, company=None) -> str:
    """One report as Markdown: bold place label, then the text, with any search
    term highlighted; JSA's own reports and seed plots carry a tag (the rest are
    Ag Trader Talk's)."""
    label, body = label_and_body(raw, location, state)
    head = f"**{_highlight(label, query)}** " if label else ""
    tag = source_tag(source, company)
    color = "blue" if source == "JSA" else "orange"
    return head + _highlight(body, query) + (f" :{color}-badge[{md_escape(tag)}]" if tag else "")


def md_table(rows: pd.DataFrame, query: str = "", dates: bool = True, new_since=None) -> str:
    """A state's reports as a Markdown table: the report date beside each report
    (left out when nothing in view has a date), in yellow when it's on or after
    `new_since`. A Markdown table rather than st.table, whose text cells stop at
    400px and turn paragraphs into columns."""
    head = ["Reported", "Report"] if dates else ["Report"]
    out = ["| " + " | ".join(head) + " |", "|" + " --- |" * len(head)]
    for r in rows.itertuples():
        day = day_label(r.date_reported)
        mark = "yellow-background" if is_new(r.date_reported, new_since) else "gray"
        cells = [f":{mark}[{day}]" if day else ""] if dates else []
        cells.append(md_report(r.raw_text, r.location, r.state, query,
                               getattr(r, "source", None), getattr(r, "source_file", None)))
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out)


# --- HTML (the PDF) -----------------------------------------------------------
_PDF_CSS = """
body { font-family: sans-serif; font-size: 10pt; color: #32373c; }
h1 { font-size: 18pt; margin: 0 0 3pt 0; }
p.sub { font-size: 9pt; color: #6b7280; margin: 0 0 8pt 0; }
h2 { font-size: 14pt; color: #0693e3; margin: 14pt 0 2pt 0; }
h3 { font-size: 11pt; margin: 10pt 0 3pt 0; }
p.r { margin: 0 0 4pt 0; }
span.d { color: #6b7280; }
"""


def report_html(df: pd.DataFrame, title: str, subtitle: str, new_since=None) -> str:
    """The PDF's HTML; a report dated on or after `new_since` on yellow."""
    e = html.escape
    out = [f"<h1>{e(title)}</h1>", f"<p class='sub'>{e(subtitle)}</p>"]
    for crop, states in sections(df):
        out.append(f"<h2>{e(crop)}</h2>")
        for name, rows in states:
            out.append(f"<h3>{e(name)}</h3>")
            for r in rows.itertuples():
                label, body = label_and_body(r.raw_text, r.location, r.state)
                day = day_label(r.date_reported)
                new = is_new(r.date_reported, new_since)
                out.append((f"<p class='r' style='background-color:{NEW_BG}'>" if new else "<p class='r'>")
                           + (f"<span class='d'>{day}&nbsp;&nbsp;</span>" if day else "")
                           + (f"<b>{e(label)}</b> " if label else "")
                           + f"{e(body)}"
                           + (f" <span class='d'>({e(tag)})</span>" if (tag := source_tag(
                               getattr(r, "source", None), getattr(r, "source_file", None)))
                              else "")
                           + "</p>")
    return "".join(out)


def to_pdf(df: pd.DataFrame, title: str, subtitle: str, new_since=None) -> bytes:
    """Letter pages, same layout as the page, numbered; reports dated on or after
    `new_since` on yellow."""
    import fitz
    story = fitz.Story(html=report_html(df, title, subtitle, new_since), user_css=_PDF_CSS)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    page = fitz.paper_rect("letter")
    frame = page + (54, 50, -54, -58)
    more = True
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(frame)
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = fitz.open("pdf", buf.getvalue())
    gray = (0.42, 0.45, 0.5)
    for i, pg in enumerate(doc, 1):
        foot = f"Page {i} of {doc.page_count}"
        x = page.width - 54 - fitz.get_text_length(foot, fontsize=8)
        pg.insert_text((x, page.height - 30), foot, fontsize=8, color=gray)
    try:
        doc.subset_fonts()          # the embedded fonts are most of the size
    except Exception:
        pass
    return doc.tobytes(garbage=3, deflate=True)


# --- Word ---------------------------------------------------------------------
def to_docx(df: pd.DataFrame, title: str, subtitle: str, new_since=None) -> bytes:
    """The same layout in Word; reports dated on or after `new_since` highlighted."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Inches, Pt, RGBColor

    gray, blue = RGBColor(0x6B, 0x72, 0x80), RGBColor(0x06, 0x93, 0xE3)
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.8)
        sec.top_margin = sec.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].font.color.rgb = blue

    doc.add_heading(title, level=0)
    sub = doc.add_paragraph().add_run(subtitle)
    sub.font.size, sub.font.color.rgb = Pt(9), gray
    for crop, states in sections(df):
        doc.add_heading(crop, level=1)
        for name, rows in states:
            doc.add_heading(name, level=2)
            for r in rows.itertuples():
                label, body = label_and_body(r.raw_text, r.location, r.state)
                p = doc.add_paragraph()
                day = day_label(r.date_reported)
                if day:
                    p.add_run(day + "   ").font.color.rgb = gray
                if label:
                    p.add_run(label + " ").bold = True
                p.add_run(body)
                tag = source_tag(getattr(r, "source", None), getattr(r, "source_file", None))
                if tag:
                    p.add_run(f" ({tag})").font.color.rgb = gray
                if is_new(r.date_reported, new_since):
                    for run in p.runs:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
